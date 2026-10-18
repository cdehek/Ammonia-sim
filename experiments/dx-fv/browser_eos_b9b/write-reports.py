"""Render reviewed machine evidence. No simulation or external tracker writes."""
import json,gzip,csv,statistics,shutil,hashlib,collections
from pathlib import Path
B=Path(__file__).resolve().parent;E=B/'evidence'
def read(n):
 p=E/n;return json.loads(gzip.decompress(p.read_bytes()) if n.endswith('.gz') else p.read_text())
def write(n,s):(B/n).write_text(s.strip()+'\n')
P=read('performance-summary.json');W=read('nonlinear-census.json');C=read('cpu-attribution.json');S=read('scientific-parity.json');runtime=read('runtime.json')
assert S['passed'] and read('adaptive-proposal-reconstruction.json')['passed'] and read('observational-removal-audit.json')['passed']
def fmt(x):return f'{x:.2f}'
def md(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(map(str,r))+' |' for r in rows])
table=[];timer=[];overhead=[];passive=[]
for engine,fixtures in P.items():
 for f,v in fixtures.items():
  table.append([engine,f,*[f"{v[m]['advanceSumMS']['median']:.2f} [{v[m]['advanceSumMS']['min']:.2f}–{v[m]['advanceSumMS']['max']:.2f}]" for m in ['original','pass-through','instrumented']]])
  timer.append([engine,f,*[fmt(v['inclusiveCostMS'][k]['median']) for k in ['NewtonInclusive','residualInclusive','engine.ph','engine.sat','engine.stateMVU','model.evaluate','model.blend','linearSolve']]])
  ov=v['overhead'];overhead.append([engine,f,fmt(statistics.median(x['passThroughDeltaMS'] for x in ov)),fmt(statistics.median(x['instrumentedDeltaMS'] for x in ov)),f"{min(x['instrumentedRatio'] for x in ov):.3f}–{max(x['instrumentedRatio'] for x in ov):.3f}"])
  passive.append([engine,f,' / '.join(fmt(x['advanceSumMS'])+(' sampled' if x['passive'] else ' control') for x in v['passive'])])
perfTable=md(['Engine','Fixture','Original ms, n=6','Pass-through ms, n=3','Diagnostic ms, n=3'],table)
timerTable=md(['Engine','Fixture','Newton','Residual','PH','SAT public','MVU','Rates','Blend','Linear + diagnostics'],timer)
overheadTable=md(['Engine','Fixture','Pass-through Δ ms','Diagnostic Δ ms','Diagnostic / bracket range'],overhead)
counts=[]
for label,key in [('Adaptive attempts','attempts'),('Accuracy rejections','rejected'),('TR-BDF2 trial integrations','TRBDF2Trials'),('Newton stages / initial residuals','stages'),('Newton corrections / Jacobians','NewtonCorrections'),('Residual evaluations','evaluations'),('Pressure FD probes',None),('Enthalpy FD probes',None),('Thermal FD probes',None),('Line-search candidates','lineSearch'),('Dense solves','linearSolves'),('Accepted half-steps','acceptedHalfSteps')]:
 vals=[]
 for f in ['cold','warm']:
  a=W['chromium'][f]['aggregate'];vals.append(a[key] if key else a['FDByKind'][{'Pressure FD probes':'pressure','Enthalpy FD probes':'enthalpy','Thermal FD probes':'thermal-energy'}[label]])
 counts.append([label,*vals])
countTable=md(['Work over .001 s','Cold','Warm'],counts)
write('NONLINEAR_WORK_CENSUS.md',f'''# Exact nonlinear-work census

Both browsers have the following identical counts for the unchanged frozen
startup workload. `evidence/nonlinear-census.json` contains per-engine, per-advance,
per-attempt, per-solve and per-TR-BDF2-trial data. Full original FD and line-search
events remain in the raw receipts. No work is inferred from elapsed time.

{countTable}

Each attempted adaptive interval performs a coarse integration and two half
integrations. Each TR-BDF2 integration has a trapezoidal stage (weight .5) and a
final BDF2 stage (weight (1−γ)/(2−γ), γ=2−sqrt(2)). All 54 / 120 Newton stages
converge. Their initial guesses are original previous PH coordinates and thermal
energy increments; the second stage uses the first stage's original result.

Cold: 54 + 2,730 + 130 = **2,914** residual evaluations.
Warm: 120 + 7,245 + 345 = **7,710**. Each correction uses 21 FD columns and one
line-search candidate. There are zero branch-rejected or failed FD probes, zero
negative probes in these fixtures, and zero rejected line-search candidates.
All accepted factors are 1. This observation does not authorize removing signed
alternatives, shrink attempts or line-search checks from other trajectories.

The original exact thermal-Jacobian option is **not enabled**. The eleven thermal
columns remain finite differences, giving 1,430 / 3,795 thermal residual probes.
No alternative setting was run. No accepted scientific configuration is replaced.

There are 8 / 17 accepted interval installations and 16 / 34 accepted half-steps.
The accuracy-rejected intervals discard 2 / 6 successfully solved half-step
candidates. All 9 / 20 coarse integrations are error-estimator references, not
accepted inventory updates. They are required numerical work under this method.
The rejected attempts consume 22 / 69 corrections and 490 / 1,536 residuals.
Accepted interval attempts consume 108 / 276 corrections and 2,424 / 6,174 residuals.
There is no natural Newton-budget exhaustion or property refusal in either fixture.

`model.trial` source events identify every trial's original start/duration,
inclusive timing, property vector, cache requests/hits/misses and child stage IDs.
`model.solveStage` events identify every stage and its recovery/cache work; original
solve IDs reset within each adaptive attempt. `model.commit` records every half
commit and `state-install` the detached atomic install. Observer arrays preserve
all outcomes and exact trial/error histories. Stage events and original solves
are ordered identically and reconcile to the totals above.

Independent M/U/V recovery occurs five times during initialization, five per
TR-BDF2 trial, five after every converged stage and five per original record:
5 + 27×5 + 54×5 + 7×5 = **445** cold;
5 + 60×5 + 120×5 + 7×5 = **940** warm. No inverse recovery is hidden inside the
Newton residual itself. All recoveries finish their pressure warm start at the
first density evaluation in these fixtures; no pressure scan/bracket fallback
is needed. Warm density evaluation still performs its original 32-step inverse.

Seven original records are generated: initial, three mandatory advance records,
and three original snapshot records. Together they contribute 56 PH / 91 public
SAT / 35 MVU calls. Additional snapshots, cloning, hashes and transport are outside
the reported advance sums. The extra three records explain the 24/39/15 count
difference from B9A's earlier component-only attribution runs.

Per-advance requests and actual accepted increments are .0001/.0004/.0005 s,
ending exactly at .0001/.0005/.001 s. Original incoming/outgoing step proposals
are in the raw timing rows. `adaptive-proposal-reconstruction.json` reconstructs
every internal unclipped proposal using the unchanged Time.factor in the SAME
browser, checks every clipped attempted dt and final nextStepSeconds exactly,
and verifies that no nonlinear-failure ceiling is active. This is post-processing
of saved errors, with zero new model/property evaluations.

See `PROPERTY_EVALUATION_AUDIT.md` for requests versus actual calls, and
`JACOBIAN_AND_LINE_SEARCH_COST.md` for cost/overlap limitations.
''')
propertyRows=[]
for label,values in [('Section PH requests',[14570,38550]),('Existing stage-cache hits',[12398,32871]),('Existing stage-cache misses / real section PH',[2172,5679]),('All real public engine PH',[8476,22072]),('Public SAT',[9261,23752]),('All lexical SAT including public',[27548,69776]),('Nested SAT beyond public',[18287,46024]),('Independent MVU / DU',[445,940]),('Single-phase PH inverse',[2976,14234]),('Mixture PH inverse',[5500,7838]),('Branch interpolation calls',[119040,601320])]:propertyRows.append([label,*values])
write('PROPERTY_EVALUATION_AUDIT.md',f'''# Accepted table-property and cache audit

The engine is root `engine.js` / `properties.json`, Git-verified against A2.
It never invokes B7, CoolProp, WASM or raw CoolProp DP. B9A's saved
`scientific-audit.json` and original baseline receipts agree with the reproduced
public counts below in both engines. These cover the full bounded workload,
including initialization and seven records; they are not one Newton iteration.

{md(['Operation','Cold','Warm'],propertyRows)}

`thermo.ph` checks finite p/h and .3–35 bar support, then calls `engine.ph` and
public `engine.sat`. It adds wet/superheat indications to a new object.
`engine.ph` delegates to lexical `inverse(p,h,'h')`. That routine evaluates
saturation and tests the original phase interval. A mixture uses statePX and
direct mixture formulas; a single-phase state uses a bounded 38-iteration
temperature-offset bisection. Each single-phase PH requires 40 branch calls:
edge check, 38 bisections, final statePT interpolation. No derivative is requested
from an EOS; the outer nonlinear solver supplies signed FD derivatives.

`sat` brackets log pressure and interpolates table saturation fields, including
log-density interpolation. `branch` brackets offset and pressure, makes two
within-pressure rows and a final interpolated row via mixRow. That means repeated
log/exp arithmetic, small temporary arrays, bracket pairs and result objects;
“table-backed” does not mean a PH calculation is a single direct array lookup.
Support checks retain the exact original epsilon, axes and missing-liquid-row
behavior. This audit changes none of them.

`inverse` computes sat(p), then statePX or statePT computes it again at the SAME
p. Those lexical calls do not go through the public proxy. Public SAT timing
therefore excludes them, while PH timing includes them. Internal counters prove
the duplication; they do not independently time its private CPU cost.

`stateMVU` validates stored mass/volume/energy, computes rho=M/V and u=U/M,
and calls lexical stateDU. stateDU evaluates atDensity at the supplied pressure
guess, normally with up to ten bounded pressure iterations and a scan/48-step
bracket fallback. Here atDensity count equals DU count: every recovery finishes
on its first pressure evaluation. AtDensity uses direct mixture relations or
a 32-iteration density/temperature inverse. Thus the cold branch count is
2,976×40=119,040; warm is 14,234×40+940×34=601,320. Recovery is independent
scientific validation, not a substitute for a failed PH result.

SAT reconciliation is exact: cold 9,261 public + 2×8,476 PH lexical + 3×445 MVU
lexical = 27,548. Warm 23,752 + 2×22,072 + 2×940 = 69,776. Cold's extra MVU
SAT is the mixture liquid-volume-fraction calculation. Public thermo.recover
also requests SAT and constructs wet/superheat fields.

## Existing PH Map

One Map is created inside each solveStage, before Newton. Its key is precisely
`${{x[2*i]}}:${{x[2*i+1]}}`, using JavaScript Number string conversion without
rounding. Section index is deliberately absent: identical PH coordinates in
different cells share a property state in that stage. Positive pressure and
nonzero fixture enthalpies avoid the +0/−0 string identity exception here; this
is not a proof for all arbitrary inputs.

Each of five section requests checks has/set/get. A successful miss stores the
genuine thermo.ph result. A thrown evaluation is not installed. Entries survive
Newton iterations and rejected line-search candidates within that stage; there
is no eviction. The Map is discarded after stage return/failure, and never
shared with the next TR-BDF2 stage, another trial, record or model session.
No numerical state or input coordinate is changed by the cache.

The measured hit rates are 85.09% cold and 85.27% warm. Thermal probes alone
produce 7,150 / 18,975 section cache hits: all five PH coordinates remain fixed.
Pressure/enthalpy probes reuse other unchanged cells; original guesses and
repeated exact coordinates can also hit. This reuse already exists in B9/B9A:
the public 8,476 / 22,072 PH counts are AFTER the stage cache.

Across the whole workload, logical PH state requests including bypassed section
hits are 20,874 / 54,943; genuine engine PH calls are 8,476 / 22,072. The section
request vector is separately 14,570 / 38,550, of which only 2,172 / 5,679 miss.
Do not compare one of those denominators with a different scope.

The two imposed boundary PH states bypass this Map and are reevaluated by every
rate evaluation. Residuals alone request them 5,828 / 15,420 times. Extra rate
evaluations at trial start and during recording add more. This exact fixed-input
repetition is a potential future opportunity, subject to deterministic-table,
validation, mutability and failure-semantic proof. No new cache is installed.

`propertyByContext` and per-stage/trial deltas retain the complete split among
initialization, section residual misses, boundary rates, coordinates, canonical
recovery, commit/event checking and recording. Property-sequence hashes prove
that the observational copy preserves every original public call's order,
arguments, return bits and thrown diagnostic meaning. Cache counters and nested
calls are read-only additions. Counts alone establish no CPU percentage.
''')
write('JACOBIAN_AND_LINE_SEARCH_COST.md',f'''# Jacobian, residual, thermal and linear-algebra cost

The 21 coordinates are interleaved five p/h pairs followed by eleven thermal
energy increments. Residual rows are five normalized mass/energy pairs plus
eleven thermal closures. Pressure/enthalpy epsilon and signed shrink loops are
unchanged. The default FD thermal setting remains in force.

All 130 / 345 Jacobian assemblies fill all 21 columns: 650/1,725 pressure,
650/1,725 enthalpy and 1,430/3,795 thermal probes. Every derivative is used on
the first positive probe here. All 130 / 345 line searches accept factor 1 on
the first candidate. A line-search residual is never counted as Jacobian work.

The following are medians of three diagnostic ADVANCE sums in ms. They are
**inclusive, nested, instrumented elapsed intervals**, not an additive CPU budget:

{timerTable}

Rates include boundary PH, temperature/domain checks, flow/donor calculations,
tube/refrigerant heat, 15 tube/fin/air links and thermal rate assembly. Residuals
also contain cached section lookups, inventory construction, energy increments,
blended quadrature and normalized balance rows. Blend repeatedly constructs
small arrays and reduces the original stage weights in the original order.
Newton also contains matrix allocation/assembly, derivative differences,
norms, original diagnostic events, linear solution and correction/line search.
Public timers do not independently split every one of these private expressions.

`linearSolve` includes the original solve PLUS read-only matrix norms,
diagonal-range and backward-residual diagnostics in the wrapper. It is therefore
an upper containing interval in the diagnostic run, not pure elimination time.
Those diagnostics do not calculate a condition number: recorded `conditionNumber`
is null. Diagonal ranges are not condition numbers. No claim that warm matrices
are poorly conditioned is supported. All observed systems are 21×21 and solve.

The original Chromium sampler independently observes unchanged solveLinear:
{C['cpu-chromium-cold']['leafFunctionWeightedMS'].get('implicit.cjs:solveLinear',0):.2f} ms cold and {C['cpu-chromium-warm']['leafFunctionWeightedMS'].get('implicit.cjs:solveLinear',0):.2f} ms warm sample-delta exposure at the solveLinear leaf;
all implicit.cjs leaves total {C['cpu-chromium-cold']['leafSourceWeightedMS'].get('implicit.cjs',0):.2f} / {C['cpu-chromium-warm']['leafSourceWeightedMS'].get('implicit.cjs',0):.2f} ms within {C['cpu-chromium-cold']['advanceSampleWeightedMS']:.2f} / {C['cpu-chromium-warm']['advanceSampleWeightedMS']:.2f} ms of sampled
advance exposure. Sampling is sparse and does not give exact CPU-cycle durations,
but it excludes the dense solve as the dominant measured browser cost here.
WebKit's containing timed intervals are also small relative to residual/Newton
intervals, with timer quantization and diagnostic overhead explicitly retained.

Each synchronous timer frame tracks measured child duration. The raw
`advanceExclusiveMS` subtracts directly nested timer intervals from the parent;
it remains an **instrumented wrapper remainder**, including unmeasured child
work, timer bookkeeping and allocation. For example Newton's remainder includes
matrix assembly, corrections and original diagnostics, and is not a pure
diagnostic-generation measurement. Rates' remainder combines physical arithmetic
and allocations. Exclusive wrapper remainders are not intrinsic CPU estimates.

Never add Newton + residual + PH + rates: each can contain another. Never
subtract a detached-copy microbenchmark as if it were an in-solver component.
Only the sampler's mutually exclusive leaf/source categories form a partition
of its explicitly included samples; unrooted GC/runtime samples cannot be
assigned to an advance and are excluded. No 100% whole-CPU pie is manufactured.

Required numerical work includes the observed finite differences, canonical
recovery, accuracy error estimation, conservation/domain checks and rejected
intervals. Potential redundancy concerns identical intermediates and temporary
organization, never skipping a branch-sensitive probe or changing tolerances.
''')
normRows=[]
for name,key in [('Attempts','attempts'),('Stages','stages'),('Newton corrections','NewtonCorrections'),('Residuals','evaluations'),('Real PH','ph'),('Branch interpolations','branch')]:
 values=[]
 for f in ['cold','warm']:
  a=W['chromium'][f]['aggregate'];n=a['publicProperties']['ph'] if key=='ph' else a['internalCounts']['internal.branch'] if key=='branch' else a[key]
  values.append(n)
 normRows.append([name,*values,fmt(values[1]/values[0])])
write('COLD_VS_WARM_DIFFICULTY.md',f'''# Why warm requires more work

The fixtures change thermodynamic inventory and hydraulic operating conditions,
as well as metal/air temperatures. Warm is not simply cold with hotter metal.

Cold begins at 3.5 bar, 500 kJ/kg: mixture quality .1397316, T≈267.811544 K,
rho≈19.84932 kg/m³, mass≈.01190959 kg per cell. Tube/fin start at 273.15 K,
air at 278.15 K. Warm begins at 3.65 bar, 1650 kJ/kg: superheated vapor,
T≈287.567822 K, rho≈2.7242515 kg/m³, mass≈.00163455 kg per cell. Tube/fin/air
start at 293.15 K. Fixed five-cell volumes and imposed inlet 4 bar/500 kJ/kg
and outlet 3.5 bar/1700 kJ/kg are unchanged. Original config and full initial
records/fluxes are retained in the baseline and census receipts.

Initial tube-to-refrigerant heat is similar: .640615 kW/cell cold versus
.669861 warm. Cold has nonzero air-to-metal links, warm initially has zero
temperature difference on those links. It is therefore unsupported to blame
warm's extra work on a greatly larger initial metal/refrigerant heat load.

Cold initial outlet flow is zero; warm initial outlet flow is .05216574 kg/s
with ≈86.0735 kW enthalpy transport. Warm's last-cell initial net mass rate is
−.05216574 kg/s and energy rate ≈−85.4036 kW. Its cell inventory is about 7.29×
smaller. These saved rates demonstrate a different rapid hydraulic startup.
They are context for the observed accuracy refinement, not an isolated causal
experiment changing one parameter.

The first .0001 s trial has normalized temporal error 1.11220 cold and 23.96491
warm. Cold rejects once, then accepts .00008 s (error .80937). Warm rejects
.0001, .00002948223 and .00001571608 s, then accepts .00001146767 s (error
.98036). All stages converge, all FD probes remain on branch and every line
search accepts immediately. The original failure-aware nonlinear ceiling stays
inactive: no natural Newton nonconvergence occurs.

{md(['Work','Cold','Warm','Warm/cold'],normRows)}

There are 2.222× as many attempted stages, and corrections/stage increase from
2.4074 to 2.8750. Cold iteration distribution is 40 solves at 2, ten at 3,
four at 5. Warm: 41 at 2, 61 at 3, 13 at 4, three at 5, one at 6 and one at 7.
Thus work amplification reflects both more attempts and more iterations/stage.
It is not a line-search rejection storm, FD branch rejection or failed recovery.

Both requested and accepted horizon are .001 s. Divide totals by .001 for
work per requested or accepted second; neither represents sustained operation.
Residuals per stage are 53.963 / 64.250; per accepted half-step 182.125 / 226.765;
per Newton correction 22.415 / 22.348. Real PH per residual is 2.9087 / 2.8628,
so warm does not have more public PH calls per residual. Instead it has more
residuals AND more expensive branches: single-phase PH fraction rises from
35.11% to 64.49%, and total branch interpolation calls rise 5.05×. Warm canonical
density recovery also uses its original 32-iteration single-phase inversion.

Advance-only PH counts are 8,437 / 22,033 (public SAT 9,197 / 23,688; MVU
420 / 915). Initialization and extra snapshots are excluded from those counts.
For wall time per residual or PH, divide a matched run's advance sum by its
advance-only count, as a mixed workload normalization, not isolated function
latency. Full-workload PH/residual ratios above intentionally include the frozen
initialization and extra records, and are work-count ratios only. Timed PH wrappers and original Chromium stacks independently identify
the property engine as substantial. Node differs in warm numerical work and
must not be pooled with the browser denominator.

No condition-number estimate or isolated temporal-error-coordinate breakdown is
available. Conditioning, compiler choices and detailed error-term causation
remain uncertain. The supported explanation is accuracy-driven step refinement,
more Newton iterations and a more expensive single-phase table workload; the
initial hydraulic/inventory evidence explains why temperature alone is inadequate.
''')
write('INSTRUMENTATION_OVERHEAD.md',f'''# Overhead and validity

Three matched rounds per engine/fixture run original → pass-through → diagnostic
→ original, each a fresh worker/loader/model with original reporting cadence.
No physics workloads or compiler jobs overlap. The complete original state,
records, attempts and property vector agree in every run; dedicated validation
also hashes every public property's input/output/throw sequence exactly.

{perfTable}

Numbers are three-advance sums for .001 simulated seconds; median [min–max].
There are six original and three runs of each wrapper mode per fixture/engine.
For each round, overhead uses the mean of its original before/after sums:

{overheadTable}

Δ columns are medians over three paired differences. Negative pass-through
deltas are observed variance/JIT/representation effects, not a speedup claim.
Exported pass-through wrapper indirection and JIT/scheduling effects can affect
this workload; their observed deviations remain explicit. Observed diagnostic differences range approximately {100*(min(x['instrumentedRatio'] for v in P.values() for f in v.values() for x in f['overhead'])-1):.1f}–{100*(max(x['instrumentedRatio'] for v in P.values() for f in v.values() for x in f['overhead'])-1):.1f}% across the brackets. A negative difference is timing variability, not negative
instrumentation cost. The diagnostic mode includes fine timers,
context strings, an additional read-only Map.has, internal counters and matrix
readouts. Do not derive intrinsic performance from these diagnostic durations
or scale all components by one global overhead ratio.

Original before/after differences are material; all raw samples remain visible
in performance-summary.json. Fresh worker does not mean cold OS cache or a
steady-state JIT. The first round is retained. Three rounds and sparse samples
do not support population confidence intervals or stable platform rankings.

The best observed clock increments are about .1 ms Chromium / 1 ms WebKit;
some bounded resolution probes observe no tick and report null. Many fine
intervals quantize to zero. Zero recorded ms does not establish zero CPU cost.
Nested elapsed timers overlap; exclusive wrapper remainders remain instrumented.

Passive Chromium controls (profiler disabled) / sampled / control advance sums:

{md(['Engine','Fixture','Before / passive / after ms'],passive)}

Only one sampled session per Chromium fixture was predefined. The sampled/control differences and their bracketing variance do not establish
an exact sampler-overhead estimate. Controls retain debugger attachment/resume but
do not enable Profiler. The separate original matrix has no CDP attachment.
WebKit has no supported sampler in this harness; its row is an original control,
not passive CPU evidence. Node inspector is supplementary and separately saved.

The initial passive harness error and a subsequently discovered sampled-control
issue are preserved under failed/superseded filenames. They are excluded from
all final statistics. Corrected final passive receipts distinguish sampled and
unsampled controls explicitly. No scientific outcomes differed in superseded
successful runs; no favorable sample selection was used.
''')
sampleRows=[]
for f in ['cold','warm']:
 c=C['cpu-chromium-'+f];sampleRows.append([f,c['advanceSamples'],c['propertyEngineStackSamples'],fmt(c['advanceSampleWeightedMS']),fmt(c['propertyEngineStackWeightedMS']),f"{100*c['propertyEngineStackFraction']:.1f}%"])
write('BROWSER_PERFORMANCE_MATRIX.md',f'''# Actual browser performance and profiler limitations

Chromium **151.0.7922.34**, desktop WPE WebKit **26.5**, Playwright **1.62.0**;
Node **24.19.0** is an additional reference. The installed system Chromium 154
is not the measured browser. Launch paths, actual user agents, versions, binary
hashes, Node/V8/Python versions and protocol identity are retained in runtime.json
and passive-chromium.json.gz. No physical Safari/iPadOS inference is made.

{perfTable}

Original timing is the authoritative unchanged advance computation boundary,
including trial/retry/acceptance/commit and mandatory record. It excludes extra
endpoint snapshots, detached state/attempt copies, exact hashes, transport and
browser startup. Initial/model times and per-command values remain in raw rows.
The harness's initializationMS begins after loader/table/engine creation and
includes model construction/create/initial record; it is not complete startup.
The unchanged original B9 worker service is separately reproduced in baseline
receipts, preserving import/factory/construction/service/RTT boundaries.

The research worker imports both original and unused diagnostic bundles. Thus
its startup/RTT also includes research module loading, large evidence copying,
bit tagging/hash work and delivery, and must not be sold as operator service
time or used to infer compilation-only cost. Worker model-advance timing and
original B9 service boundaries remain distinct. No startup optimization is tested.

Original detached state/attempt copy medians are {P['chromium']['cold']['original']['stateAndAttemptCopyMS']['median']:.2f}/{P['chromium']['warm']['original']['stateAndAttemptCopyMS']['median']:.2f} ms cold/warm Chromium
and {P['webkit']['cold']['original']['stateAndAttemptCopyMS']['median']:.2f}/{P['webkit']['warm']['original']['stateAndAttemptCopyMS']['median']:.2f} ms WebKit over the three endpoint copies. Extra original snapshot
record sums are {P['chromium']['cold']['original']['extraSnapshotRecordMS']['median']:.2f}/{P['chromium']['warm']['original']['extraSnapshotRecordMS']['median']:.2f} ms Chromium and {P['webkit']['cold']['original']['extraSnapshotRecordMS']['median']:.2f}/{P['webkit']['warm']['original']['extraSnapshotRecordMS']['median']:.2f} ms WebKit at available resolution.
These are actual outside-advance boundaries, not deductions from physical CPU.
Original diagnostic generation remains inside the solver and is not isolated as
a complete independent component. Display transport is not rebenchmarked here;
B9A's preserved evidence remains the secondary transport/copy reference.

## Original-source sampling

Chromium CDP Profiler requests 500 µs sampling on the dedicated original worker.
Raw cpu-chromium-*.json preserves call frames, source positions, samples and
timeDeltas. Only stacks containing accepted-bundle's advance function are used
for advance attribution; hashing, serialization and module startup are excluded.
The source range map is derived from exact accepted slices, not function names
alone. Anonymous frames retain their source file/line identity.

{md(['Fixture','Advance samples','Engine stack samples','Weighted advance ms','Weighted engine ms','Engine weighted fraction'],sampleRows)}

For cold, coupled-model leaves total {C['cpu-chromium-cold']['leafSourceWeightedMS'].get('coupled-model.cjs',0):.2f} ms, table engine {C['cpu-chromium-cold']['leafSourceWeightedMS'].get('engine.js',0):.2f} ms, coupled
Newton {C['cpu-chromium-cold']['leafSourceWeightedMS'].get('coupled-implicit.cjs',0):.2f} ms and linear module {C['cpu-chromium-cold']['leafSourceWeightedMS'].get('implicit.cjs',0):.2f} ms of sampled exposure. Warm: table engine
{C['cpu-chromium-warm']['leafSourceWeightedMS'].get('engine.js',0):.2f} ms, coupled-model {C['cpu-chromium-warm']['leafSourceWeightedMS'].get('coupled-model.cjs',0):.2f} ms, coupled Newton {C['cpu-chromium-warm']['leafSourceWeightedMS'].get('coupled-implicit.cjs',0):.2f} ms and linear {C['cpu-chromium-warm']['leafSourceWeightedMS'].get('implicit.cjs',0):.2f} ms.
MixRow/branch interpolation and the model's residual/quadrature arrays are visible
hot paths. The property stack is substantial but does not account for every
millisecond. The main supported ranking is **nonlinear residual/property work**,
then smaller linear/record/commit work, with separate diagnostic copying expense.

These are sample-delta-weighted stack exposure estimates, not exact exclusive
CPU cycles. Browser scheduling and lost/unrooted runtime/GC samples matter.
The cold profile has only {C['cpu-chromium-cold']['advanceSamples']} included samples; one profile per fixture cannot
establish stable percentages. No total-CPU 100% decomposition is claimed.

WebKit's Playwright `new_cdp_session` returns “CDP session is only available in
Chromium”. This establishes a limitation of the available automation interface,
not absence of every WebKit inspector/profiler in principle. No unsupported raw
inspector protocol is substituted. WebKit uses exact work counters and direct
instrumented boundary timings, with explicit overhead; no WebKit sampled stack
or private-function CPU percentage is asserted.

## Engine comparison and uncertainty

Both browsers execute identical operation counts and match their own saved
scientific baseline exactly. Their internal floating values are not assumed
cross-engine identical. Original medians differ, with overlapping scheduling/JIT
variability and coarse WebKit timing. Warm PH inclusive diagnostic medians are
{P['chromium']['warm']['inclusiveCostMS']['engine.ph']['median']:.2f} ms Chromium / {P['webkit']['warm']['inclusiveCostMS']['engine.ph']['median']:.2f} ms WebKit; those values include observational counter
overhead and cannot establish a pure intrinsic property-engine ratio.

Node's warm reference uses 22,013 PH / 23,693 SAT / 940 MVU, one fewer Newton
correction than these browsers; cold matches their public count vector. Its
original, isolated-copy and direct-CommonJS references agree within Node, including
full original property sequences. Node timings/samples remain separate. Neither
Node nor desktop WebKit stands in for physical iPadOS Safari.

The full bounded workload is tiny startup only. No longer-duration throughput,
real-time capability, cooling-equipment calibration, speedup or device ranking
is extrapolated. All samples, significant spreads, profiler failures and limits
remain in the evidence rather than being replaced with favorable observations.
''')
write('SCIENTIFIC_CONTRACTS_AND_BLOCKERS.md','''# Preserved scientific contracts and unresolved qualification

Accepted A2 is the original five-volume, 21-coordinate table-backed model.
No EOS equation, coefficient, property input, cache decision, Newton/TR-BDF2
operation, signed FD, thermal Jacobian setting, tolerance, growth proposal,
domain/phase rule, conservation budget or acceptance policy changes in B9B.

Preserved contracts: canonical recovery p≤.1 Pa, h≤.1 J/kg, T≤1e-4 K;
normalized fluid mass/energy residual≤1e-11; standalone p strictly <10 Pa,
h strictly <50 J/kg, T strictly <.01 K; independent reference resolution
1e-5 Pa / 1e-8 J/kg. Pressure support .3–35 bar absolute; original table
temperature/phase support; thermal domain 200–400 K. M/U/V arithmetic and
ordering remain exact. Aggregate/section mass <1e-8 kg, energy <1e-6 kJ,
thermal nonlinear closure 1e-12 K, original Newton 30-iteration limit,
TR-BDF2, minStep 1e-9, exact accepted-state rollback all remain in force.

The browser parity receipt independently rebuilds the full exact science/observer
hash from saved B9A baseline-*-export-fixed.json.gz (not only checkpoint hashes),
and compares every original state and record field,
including M/U/V/p guesses, recovered p/T/h/rho/u/quality/phase/wet, eleven
thermal energies/temperatures, all fluxes and ledgers, clock/events/boundaries,
residuals/iteration/FD/line-search/stage/retry decisions and growth histories.
Every converged observer event independently satisfies original fluid/thermal
normalized limits; scaling its thermal norm retains ≤1e-12 K. Record budgets
and independently signed thermal ledgers pass at each checkpoint. A gauge or
compact projection is never used as authoritative state.

This bounded work preserves prior standalone/reference contracts by exact
source identity; it does not repeat or extend their complete qualification.
Four DU saturation endpoints remain unresolved. Eleven invalid DP path/input
identities and forty-four historical invalid DP observations remain. Seven
original 35-bar browser refusals per engine remain refusals. Two saved inward
false-acceptance hazards remain unresolved. Exact M/U/V and rounded-DU boundary
inputs remain distinct; neither label nor rounded pair certifies the other.
Complete phase/topology certification remains unproved. Experimental CoolProp
backend qualification, complete A2/long-duration qualification, nine-section
spatial convergence, production protections and physical Safari/iPadOS
performance remain open. No faster table-engine result qualifies CoolProp.

The inherited 35-row B8 issue-priority JSON/CSV is preserved byte-for-byte under
evidence and included in the local workbook. No external tracker was edited.
Production main, PR #3, demo release, accepted A2, previous research, UI,
Live Plant, operator workstation and Cloudflare remain outside this work.
No merge, deployment, EOS repair, new protection or production integration.
''')
write('WORK_CLASSIFICATION.md','''# Necessity versus future opportunity

Classification is about dependencies under the original algorithm. It is not
authorization to remove work. A rejected trial or probe is never labeled waste
just because it did not become the accepted inventory.

| Operation | Classification | Evidence / limitation |
|---|---|---|
| Changed section PH state / branch inverse | REQUIRED NUMERICAL WORK | Supplies density, temperature and phase for original conservative residual |
| Stage exact-key Map lookups | REQUIRED NUMERICAL WORK under existing implementation | Already avoids most section PH engine calls; no new cache here |
| Same boundary PH repeated across residuals | POTENTIALLY REDUNDANT WORK | Exact fixed spec/input within stage; purity/refusal/mutation equivalence unproved for future reuse |
| Same-operation repeated lexical saturation | POTENTIALLY REDUNDANT WORK | inverse then statePX/PT evaluate identical p; support/phase checks must survive any future reuse |
| All original Jacobian FD probes | REQUIRED NUMERICAL WORK | Current accepted fixture uses all 21 FD columns; signed/branch alternatives must remain for other states |
| Line-search candidates | REQUIRED NUMERICAL WORK | Original Newton acceptance check, even though first candidate always succeeds here |
| Coarse TR-BDF2 solution and two half solutions | REQUIRED NUMERICAL WORK | Original temporal estimator and unextrapolated accepted inventories |
| Accuracy-rejected intervals | REQUIRED NUMERICAL WORK | Actual error >1; original refinement preserves accuracy |
| Flow, donor enthalpy, heat links, thermal/fluid balances, scaling | REQUIRED NUMERICAL WORK | Defines original coupled equations and acceptance norms |
| Rebuilt thermal ID/index Map | POTENTIALLY REDUNDANT WORK | Same spec topology; individual allocation CPU cost not isolated |
| Blend arithmetic | REQUIRED NUMERICAL WORK | Original quadrature/ordering; temporary-allocation avoidability is UNCERTAIN |
| Dense 21-coordinate elimination/correction | REQUIRED NUMERICAL WORK | Original Newton linear solution, measured secondary cost |
| Dense scratch allocation | POTENTIALLY REDUNDANT WORK | Low-priority hypothesis; exact pivots/order/alias reset obligations remain |
| Canonical MVU recovery and stage agreement | REQUIRED NUMERICAL WORK | Independent inventory consistency can refuse a converged PH solve |
| SAT crossing checks / domain proof | REQUIRED NUMERICAL WORK | Original phase/event and physical-exit evidence, not optional display |
| Detached ledger preparation and atomic commit | REQUIRED NUMERICAL WORK | Required exact rollback and installation semantics |
| Original record scientific reconstruction/validation | REQUIRED work in frozen scientific/reporting contract | Can fail honestly; cannot substitute a gauge or drop cadence |
| Record formatting/copying subcomponent | UNCERTAIN CPU attribution | Containing record includes real properties/validation; no complete private split |
| Original observer iteration/FD/line-search copies | DIAGNOSTIC/PRESENTATION OVERHEAD | Source reads completed values into diagnostics only; exact frozen observer/cadence still retained |
| Outside-advance state/attempt copies, evidence hashing and delivery | DIAGNOSTIC/PRESENTATION OVERHEAD | Identified owner/transport boundaries; no feedback to state |
| Private diagnostic construction, GC and all temporary allocation costs | UNCERTAIN detailed CPU attribution | Sampling/containing timers do not independently isolate all components |
| B9B timers/counters/matrix readouts | DIAGNOSTIC overhead | Measured overhead controls; absent from original intrinsic-performance conclusions |

The dominant measured source family is nonlinear residual/property evaluation.
The evidence supports a narrow exact-input repetition hypothesis, not a claim
that convergence/accuracy/phase checks are unnecessary. All proposed changes
remain future reviewed work; no optimization is implemented in B9B.
''')
write('TIMING_OVERLAP_AND_LIMITATIONS.md','''# Timing ownership and overlap ledger

| Boundary | Contains / excludes | Valid interpretation |
|---|---|---|
| Original advance | Trials, stages, original checks/diagnostics, commits, mandatory final record | Authoritative computation elapsed control |
| Trial | Coordinates, canonical checks, initial rates and two stages | Inclusive; overlaps stages |
| Stage | Newton plus post-solve recovery/checks | Inclusive; overlaps Newton/recovery |
| Newton | Initial/FD/line-search residuals, matrix solve/assembly, correction, diagnostics | Inclusive; overlaps residual/linear |
| Residual | Stage cache, PH misses, rates, inventories, blend, normalization | Inclusive; overlaps rates/PH/blend |
| Rates | Boundary PH, temperature checks, flow/heat/thermal assembly | Inclusive; overlaps PH |
| Thermo PH | Engine PH, public saturation, guard and derived indications | Inclusive; overlaps both properties |
| Engine PH | Lexical inverse, saturation, branch interpolation | Public timer includes internal calls |
| Engine MVU | DU/atDensity inverse and any lexical SAT | Independent canonical operation; internal SAT overlaps |
| Public SAT | Exported invocations only | Does not include lexical internal SAT |
| Linear timer | Original dense solve and read-only matrix diagnostics | Containing diagnostic interval, not pure elimination |
| Wrapper exclusive remainder | Parent minus immediate measured child intervals | Instrumented remainder, not intrinsic CPU |
| Detached state/attempt copy | After original advance, before evidence hash/transport | Separate actual boundary, not part of physical advance |
| Snapshot record | Extra original endpoint record | Recovery/flux/validation, not just copying |
| Worker RTT | Imports, model work, evidence copies/hashes, dispatch/transport | Research delivery, not solver-only or compilation-only |
| Sample leaf/source | Mutually exclusive included stack samples | Sparse source attribution; excludes unrooted GC/outside-advance work |

No nested interval sums or percentage breakdown across overlapping boundaries
are valid. Timing calls and bookkeeping themselves consume time. Coarse clocks
can yield zero intervals and imprecise remainders. The raw exclusive map is kept
for audit but not advertised as a calibrated intrinsic cost model. No microbench
subtraction or massaged 100% partition is used.

Private thermal/fluid subexpressions, branch-check/normalization arithmetic,
diagnostic generation/allocation and GC are not each independently attributable.
Raw CPU source lines are observation positions, subject to JIT/inlining; they
are not proof of per-expression exclusive cost. No WebKit sampled CPU breakdown
or independently measured condition number is available. No long-run memory,
retention, sustained throughput or physical-device result is established.
''')
write('FAILURE_AND_ROLLBACK.md','''# Bounded failure/rollback receipts

PASS in actual Chromium and desktop WebKit. `failure-*.json.gz` contains the
unchanged original B9/C1/B6 scenario harness, additional original C1 sessions
with/without public property counters, and accepted-original versus isolated
observational-copy fault-path comparisons. Synthetic injections are explicitly
marked and are not natural thermodynamic trajectories.

The original C1 run proves supported nonzero exchange, transient refusal,
persistent refusal, original numerical stop and independently confirmed 200/400 K
outward thermal exits. Its additional original B9 rollback check verifies all
accepted physical bits and array/ledger/individual-cell identities at every
rejected trial; a failed accepted-state record returns no usable property state.

Additional C1 sessions compare transient, persistent inverse-recovery-failure,
boundary-unresolved, failed accepted-state record, thermal200 and thermal400.
Public pass-through/count instrumentation leaves their complete scientific
records, attempts, errors and proofs identical. Persistent synthetic scenarios
retain 20 attempts, 19 halvings, minStep 1e-9, Newton limit 30, numerical stop
and unproved physical exit. Boundary-unresolved is not relabeled as a proven
physical excursion. Original C1 recordFailure metadata is retained exactly.

The accepted-source observational copy is separately compared with accepted A2
on transient/persistent injected solver refusals, failed recording and genuine
200/400 K outward exits. It retains all original states/records/attempts/errors
and exact physical/reference rollback. These tests use the copied observational
source hooks as read-only counters; the C1 source is never transformed. They
do not claim the accepted A2 source has acquired C1's newer typed error contract.
Finite timers are qualified by full successful-fixture/property-sequence parity;
failure source hooks and exported-counter error delegation have the bounded
coverage above. No broader fault/observer-exception campaign is claimed.

Profiling neither repairs properties nor adds usable states after refusal.
The independently rechecked exact-boundary thermal proof remains separate from
refrigerant numerical refusal. Trial rollback is an original solver operation;
worker termination is disposable-session destruction and proves no rollback.
No new lifecycle implementation, original-error classification or acceptance
rule is introduced. B9A lifecycle and protocol sources/evidence remain untouched.
''')
write('ENGINEERING_REPORT.md',f'''# B9B engineering decision

**B9B PERFORMANCE FORENSICS PASS**

The original five-section computation cost is sufficiently attributed for one
narrow subsequent engineering decision. This is a diagnostic milestone, not a
faster solver, backend qualification or production integration.

| Required decision | Result |
|---|---|
| Accepted A2 source integrity | PASS — fourteen original Git identities, all historical files preserved |
| Exact scientific parity | PASS — both saved same-engine baselines; all state/record/attempt and property-sequence checks |
| Failure/rollback | PASS — original bounded C1 contract plus observational-copy fault-path parity; scope in FAILURE_AND_ROLLBACK.md |
| Work census | PASS — residual/Jacobian/FD/line-search/stage/trial/cache counts reconcile |
| Property cache attribution | PASS — requests, hits, misses, real engine and nested lexical operations distinguished |
| CPU profiling attribution | PASS — bounded original Chromium/Node sampling plus WebKit timed boundaries; private/exclusive limits retained |
| Primary bottleneck | Nonlinear residual/property evaluation: repeated table inversions/interpolation and coupled rate/residual/quadrature assembly |
| Primary bottleneck confidence | HIGH for joint residual/property dominance; MODERATE for stable subcomponent CPU shares |
| Instrumentation overhead | Observed diagnostic bracket ratios {min(x['instrumentedRatio'] for v in P.values() for f in v.values() for x in f['overhead']):.3f}–{max(x['instrumentedRatio'] for v in P.values() for f in v.values() for x in f['overhead']):.3f}; pass-through variance material |
| Smallest recommended next milestone | Separately reviewed, table-only stage-boundary PH reuse equivalence design and isolated prototype |
| Further isolated research | GO |
| Operator-workstation integration | NO-GO |

Gate 1 PASS: exact accepted closure, frozen settings/cadence, actual Chromium and
desktop WebKit match original saved science and budgets. Gate 2 PASS: every
attempt/stage/iteration/probe/line search and original cache is accounted for,
with cold/warm explanation and no invented exclusive accounting. Gate 3 PASS:
original decisions and scientific/property sequences remain identical; bounded
failure semantics and rollback remain; measured overhead/limitations are explicit.
Gate 4 PASS: actual browser timing and original sampling support dominant versus
secondary costs, four narrow future hypotheses, uncertainty and one next action.

The 21-coordinate FD problem needs 21 probes per correction. Adaptive TR-BDF2
uses coarse plus two half integrations, six stage solves per attempted interval.
Cold executes 9 attempts / 54 stages / 130 corrections / 2,914 residuals; warm
20 / 120 / 345 / 7,710. All stages converge, all line searches accept factor 1,
and all signed probes use the first positive candidate. Accuracy refinement,
not repeated Newton failure, drives the observed retries. Those retries and
coarse solutions are required under the original method.

Warm starts with a much smaller superheated inventory and substantial outlet
flow, then triggers larger temporal error and more iterations per stage. It also
requires more expensive single-phase table inverses. Branch interpolation calls
increase 119,040→601,320. The existing stage cache already avoids 12,398 / 32,871
section PH evaluations. Boundary PH states bypass it, so unchanged imposed
boundary pairs are evaluated repeatedly in every residual, including thermal
FD probes. This is the best supported potential redundancy, subject to exact
table purity, phase/validation/refusal, ownership and rollback proof.

Original .001-s advance-sum medians are {P['chromium']['cold']['original']['advanceSumMS']['median']:.2f}/{P['chromium']['warm']['original']['advanceSumMS']['median']:.2f} ms cold/warm Chromium and
{P['webkit']['cold']['original']['advanceSumMS']['median']:.2f}/{P['webkit']['warm']['original']['advanceSumMS']['median']:.2f} ms WebKit. The diagnostic copy is substantially slower; its raw
inclusive times are not intrinsic solver times. Original Chromium stacks show
both table interpolation and coupled residual construction as substantial,
with small dense-solve exposure. Outside-advance state/attempt copying remains
secondary. Complete private diagnostic generation is not independently isolated.
No property-only or solver-only speedup estimate, 100% CPU split, sustained
real-time capability, device prediction or matrix-conditioning explanation is
asserted. WebKit has no supported sampler through this harness.

No optimization hypothesis is implemented. No accepted solver/EOS/thermal/phase/
domain/tolerance/growth/budget changes, CoolProp rebuild, horizon extension, full
A2 campaign, UI/Live Plant/workstation edit, merge or deployment occurred.
All inherited scientific/backend/device blockers remain in
SCIENTIFIC_CONTRACTS_AND_BLOCKERS.md. Commit/push the isolated additive branch
only and stop for explicit engineering review. Do not automatically begin the
recommended prototype.
''')
write('MECHANICS_REPORT.md',f'''# 🔧 Ammonia Lab — The Mechanic’s Report

The virtual coil is doing far more than updating five temperature readings.
It stores refrigerant mass and internal energy in five sections, plus heat in
the tube, fins and air. Those inventories must agree with pressure, enthalpy,
density, flow and heat transfer at the same instant.

Newton tries a set of pressures, enthalpies and metal/air energies, checks all
21 balances, then nudges each coordinate to learn how the balances respond.
Every correction here requires 21 separate probe evaluations and another
evaluation to check the proposed correction. TR-BDF2 then compares a full
interval with two half intervals before accepting the half-step inventories.
That is six nonlinear stage solves per attempted interval. A rejected interval
is the model refusing an inaccurate advance. The accepted inventories stay at
their last proven values until a smaller interval passes the original numerical
test; the rejection is not automatically wasted calculation.

The property engine also does real work. For a two-phase mixture it can use
saturation and mixture relations. For superheated refrigerant it searches the
bounded property table 38 times to recover temperature from pressure and
enthalpy. Each search interpolates several rows and evaluates density logarithms
and exponentials. A table lookup is therefore often an inverse calculation.

Warm begins with superheated vapor and much less mass in each section. It also
starts discharging through the outlet while cold initially has zero outlet
flow. The initial tube/refrigerant heat transfer is quite similar in the two
cases. The trace shows warm needs smaller accurate intervals and more Newton
corrections, and each property workload contains more expensive vapor inverses.
It is not simply “higher temperature makes the computer slow.” There are no
Newton failures or line-search rejection storms in these short fixtures.

The model already reuses exact section PH results inside each stage. Cold has
12,398 cache hits, warm 32,871. However, it calculates the two unchanged imposed
boundary states again whenever it evaluates rates. That repetition is a concrete
future opportunity to investigate. We have not removed it. A boundary result
could only be reused after proving identical values, support checks, failure
meaning and ownership for the exact table engine and exact unchanged inputs.

In the unchanged runs, advancing only .001 simulated seconds takes median
advance sums of about {P['chromium']['cold']['original']['advanceSumMS']['median']:.0f}/{P['chromium']['warm']['original']['advanceSumMS']['median']:.0f} ms cold/warm in Chromium and {P['webkit']['cold']['original']['advanceSumMS']['median']:.0f}/{P['webkit']['warm']['original']['advanceSumMS']['median']:.0f} ms in desktop
WebKit. Most measured computation belongs to nonlinear residual/property work;
the small matrix solve is not the main expense. Separately copying the full
research state and attempt history adds a separate measured cost over the endpoint
copies. We cannot honestly assign every remaining millisecond to physics versus
diagnostic construction. Fine profiling itself adds substantial overhead, so
its slower timings are not used as the original coil's performance.

Any future improvement must conserve the same refrigerant and heat, keep the
same phase and support checks, reject the same inaccurate steps and restore the
same accepted inventories exactly after failure. A fast calculation that hides
a property refusal or quietly creates/removes mass or energy is not progress.

The next engineering review should consider a very small table-only experiment:
prove whether the unchanged boundary PH calculations can be reused within one
stage. Keep Newton, signed probes, time-step acceptance and canonical recovery
intact. If that proof fails, stop and examine duplicate saturation intermediates
within a single property call. B9B provides evidence for that decision; it claims
no speedup. The responsive operator workstation, longer operation and physical
iPad/Safari performance still require separate qualification.
''')
write('EXECUTION_NOTES.md','''# Execution notes and retained unsuccessful evidence

The workspace began without a checkout. The first sandboxed clone could not
reach its proxy; explicitly network-enabled execution succeeded. The checkout
was clean at exact B9A, then the isolated B9B branch was created. No remote or
protected file was changed by setup. Gate 0 verified 2,675 historical files and
14 A2 sources against original Git objects, plus all C1/B6/B7/B9 files at their
protected commits. Pre-instrumentation design commit: 5fae85a.

Browser dependencies were installed only in /workspace/b9-runtime using the
unchanged pinned B9 setup; its warnings/log are retained. No host/operator
workstation dependency was installed. Actual measured browsers are the original
Playwright Chromium 151 and WPE WebKit 26.5, not system Chromium 154.

Original B9 service and direct frozen-worker baselines passed before diagnostic
runs. External Chromium sampling followed; the first driver expression defined
and inadvertently invoked a function without arguments. That harness failure is
retained as passive-chromium-harness-failed.json.gz and its log. It supplied no
qualified CPU profile. The corrected pass then accidentally enabled sampling on
its bracketing controls too; those successful-science receipts/profiles are
retained under superseded-sampled-controls-* and excluded from final overhead
and CPU summaries. The final bounded passive pass enables Profiler only on the
designated sampled worker and records controls separately. No scientific outcome
or property count changed in those successful superseded runs.

The diagnostic copy inserts observations only, with a full exact diff and an
independent deletion audit restoring every original byte. Node smoke checks
preceded the three browser measurement rounds. Final successful-fixture validation
checks exact scientific fields separately from profiler metadata and compares
every original public property request/result sequence. Additional bounded
failure cases use original C1 sources and original-vs-observational A2 copies.

Final review found that the first main-matrix original control unnecessarily
constructed unused context strings and a no-op timing closure on public property
calls. Those context/closure operations are now bypassed in the original and
pass-through controls. Earlier controls/matrix/CPU receipts are retained under
superseded-control-context-* and excluded from final statistics. The predefined
baseline/passive/matched matrix and Node reference were rerun after this harness
correction; this was not a search for favorable timing. All scientific data in
the earlier matrix remained unchanged. Final scientific validation/failure
checks follow the corrected matrix, sequentially. Worker dispatch later gained
only the additional isolated-copy failure command; measured profile.mjs/scientific
copy implementation was not changed during or after the matrix. A post-run
same-browser Time.factor replay reconstructs every adaptive proposal without
executing any further model or property work.

Offline analysis corrected two parser issues: sorting a null clock-resolution
probe with numeric values, and excluding its own CPU summary from the raw-profile
glob on reruns. These changed no raw evidence or experiment. Final analysis
asserts exact parity, original budgets, complete count reconciliation and
observational insertion identity. Large raw gzip receipts preserve full original
arrays rather than substituting compact/gauge snapshots.

All research files are additive under B9B. Historical file/ref/PR preservation,
final branch ancestry and clean working tree are sealed separately. No mutation
of a production ref, deployment or review message to another person is made.
''')
write('README.md','''# Research B9B — solver/property performance forensics

**B9B PERFORMANCE FORENSICS PASS.** The accepted table-backed five-section model
is unchanged. Its short startup cost is attributed to nonlinear residual/property
work, with a measured existing cache and a narrow future boundary-reuse hypothesis.
This milestone implements no numerical optimization. Operator integration NO-GO.

Start: `355db28d853817e1931295d6ec2f007876d53f90`; accepted A2:
`be9eb35a4ddf5098e685182f8104685127c6a6ee`. Branch:
`research/browser-eos-b9b-solver-property-profiling-20261017`.
The pre-instrumentation design was independently committed as `5fae85a`.

Read ENGINEERING_REPORT.md for gates/decision, MECHANICS_REPORT.md for practical
meaning, PROFILING_DESIGN.md for the frozen design, NONLINEAR_WORK_CENSUS.md,
PROPERTY_EVALUATION_AUDIT.md, JACOBIAN_AND_LINE_SEARCH_COST.md,
COLD_VS_WARM_DIFFICULTY.md, INSTRUMENTATION_OVERHEAD.md,
BROWSER_PERFORMANCE_MATRIX.md and OPTIMIZATION_HYPOTHESES.md for findings.
TIMING_OVERLAP_AND_LIMITATIONS.md, FAILURE_AND_ROLLBACK.md and
SCIENTIFIC_CONTRACTS_AND_BLOCKERS.md preserve scope and obligations.

Raw evidence contains both actual browsers, original service reproduction,
matrix/validation/failure receipts, original and nested property/cache counters,
all original attempts, complete state/records, sampled stacks and Node reference.
scientific-parity.json, nonlinear-census.json, adaptive-proposal-reconstruction.json,
cpu-attribution.json and performance-summary.json are the machine-readable results.
milestone-status.json/csv/xlsx support the Laboratory tracker; no external sheet
or operator workstation was changed. Failed/superseded receipts remain labeled.

Reproduction from repository root (use explicit process timeouts):

```sh
# Read-only source checks / deterministic observational build
python experiments/dx-fv/browser_eos_b9b/build.py
python experiments/dx-fv/browser_eos_b9b/verify-copy.py
# Browser setup writes only /workspace/b9-runtime, requires network
timeout 420 python experiments/dx-fv/browser_eos_b9a/setup-browsers.py
# Each stage refuses to overwrite existing evidence. For a rerun, move prior
# B9B outputs to a labeled B9B receipt directory; never overwrite frozen history.
PLAYWRIGHT_BROWSERS_PATH=/workspace/b9-runtime/browsers timeout 300 python experiments/dx-fv/browser_eos_b9b/run-browser.py --stage baseline
PLAYWRIGHT_BROWSERS_PATH=/workspace/b9-runtime/browsers timeout 300 python experiments/dx-fv/browser_eos_b9b/run-browser.py --stage passive
PLAYWRIGHT_BROWSERS_PATH=/workspace/b9-runtime/browsers timeout 300 python experiments/dx-fv/browser_eos_b9b/run-browser.py --stage matrix
PLAYWRIGHT_BROWSERS_PATH=/workspace/b9-runtime/browsers timeout 300 python experiments/dx-fv/browser_eos_b9b/run-browser.py --stage validation
PLAYWRIGHT_BROWSERS_PATH=/workspace/b9-runtime/browsers timeout 300 python experiments/dx-fv/browser_eos_b9b/run-browser.py --stage failure
timeout 60 node experiments/dx-fv/browser_eos_b9b/run-node.mjs
PLAYWRIGHT_BROWSERS_PATH=/workspace/b9-runtime/browsers timeout 60 python experiments/dx-fv/browser_eos_b9b/reconstruct-proposals.py
timeout 120 python experiments/dx-fv/browser_eos_b9b/analyze.py
python experiments/dx-fv/browser_eos_b9b/write-reports.py
python experiments/dx-fv/browser_eos_b9b/preservation.py
```

Network-enabled execution is needed for localhost browser servers in a sandbox
that isolates networking. Node is supplemental. No CoolProp compilation or EOS
modification is required or permitted. Stop after isolated commit/push for explicit
engineering review; do not automatically implement an optimization.
''')
# Preserve inherited issue tracker and create local flat tracker/workbook.
for ext in ['json','csv']:
 src=B.parent/f'browser_eos_b9a/evidence/inherited-B8-issue-priorities.{ext}'
 shutil.copyfile(src,E/src.name)
status=dict(milestone='B9B',classification='B9B PERFORMANCE FORENSICS PASS',acceptedA2Integrity='PASS',exactScientificParity='PASS',failureRollback='PASS',workCensus='PASS',propertyCacheAttribution='PASS',CPUProfilingAttribution='PASS',primaryBottleneck='nonlinear residual/property evaluation',confidence='HIGH joint; MODERATE subcomponent shares',instrumentationOverheadRatioRange=[min(x['instrumentedRatio'] for v in P.values() for f in v.values() for x in f['overhead']),max(x['instrumentedRatio'] for v in P.values() for f in v.values() for x in f['overhead'])],recommendedNextMilestone='Review table-only stage-local immutable boundary PH reuse equivalence design/prototype; no implementation in B9B',furtherIsolatedResearch='GO',operatorWorkstationIntegration='NO-GO',start='355db28d853817e1931295d6ec2f007876d53f90',branch='research/browser-eos-b9b-solver-property-profiling-20261017',actualHorizonSeconds=.001,optimizationImplemented=False,remainingBlockers='4 DU endpoints; 11 DP identities / 44 observations; 7 35-bar refusals per engine; 2 inward hazards; exact MUV vs rounded DU; topology; EOS; A2/long-duration; nine-section; production protections; physical Safari/iPadOS')
(B/'milestone-status.json').write_text(json.dumps(status,indent=2)+'\n')
flat=[dict(milestone='B9B',classification=status['classification'],furtherResearch='GO',operatorIntegration='NO-GO',engine=engine,fixture=f,mode=mode,horizonSeconds=.001,advanceMedianMS=v[mode]['advanceSumMS']['median'],advanceMinMS=v[mode]['advanceSumMS']['min'],advanceMaxMS=v[mode]['advanceSumMS']['max'],samples=v[mode]['advanceSumMS']['n'],PH=W[engine][f]['aggregate']['publicProperties']['ph'],SAT=W[engine][f]['aggregate']['publicProperties']['sat'],MVU=W[engine][f]['aggregate']['publicProperties']['stateMVU'],advanceOnlyPH=W[engine][f]['aggregate']['advancePublicPropertyCounts']['ph'],residuals=W[engine][f]['aggregate']['evaluations'],NewtonCorrections=W[engine][f]['aggregate']['NewtonCorrections'],stageCacheHits=W[engine][f]['aggregate']['cache']['hits'],stageCacheMisses=W[engine][f]['aggregate']['cache']['misses'],exactScientificParity=True) for engine,fixtures in P.items() for f,v in fixtures.items() for mode in ['original','pass-through','instrumented']]
normalizations=[]
for engine,fixtures in P.items():
 for fixture,v in fixtures.items():
  a=W[engine][fixture]['aggregate'];ms=v['original']['advanceSumMS']['median'];ph=a['advancePublicPropertyCounts']['ph']
  normalizations.append(dict(engine=engine,fixture=fixture,requestedSeconds=.001,acceptedSeconds=.001,advanceMedianMS=ms,residualsPerRequestedOrAcceptedSecond=a['evaluations']/.001,residualsPerAttemptedStage=a['evaluations']/a['stages'],residualsPerAcceptedHalfStep=a['evaluations']/a['acceptedHalfSteps'],residualsPerNewtonCorrection=a['evaluations']/a['NewtonCorrections'],advancePHPerResidual=ph/a['evaluations'],advanceMedianMSPerResidual=ms/a['evaluations'],advanceMedianMSPerRealPH=ms/ph,limitation='Mixed whole-advance normalization; never an isolated property/residual latency'))
(E/'work-normalizations.json').write_text(json.dumps(normalizations,indent=2)+'\n')
failureRows=[]
for engine in ['chromium','webkit']:
 for r in read('failure-'+engine+'.json.gz')['failure']['results']:
  p=r['proof'];failureRows.append(dict(engine=engine,scenario=r['scenario'],instrumentedPublicCounter=r['instrumented'],synthetic=p['synthetic'],failures=p['failures'],floor=p['floor'],maxIterations=p['maxIterations'],acceptedBitsUnchanged=p['acceptedBitsUnchanged'],identityUnchanged=p['identityUnchanged'],physicalTrajectoryClaim=p['physicalTrajectoryClaim']))
(E/'milestone-performance-data.json').write_text(json.dumps(dict(status=status,performance=flat,normalizations=normalizations,work={engine:{f:v['aggregate'] for f,v in fixtures.items()} for engine,fixtures in W.items()},failures=failureRows,runtime=runtime),indent=2)+'\n')
with (B/'milestone-status.csv').open('w') as file:
 writer=csv.DictWriter(file,fieldnames=list(flat[0]),lineterminator='\n');writer.writeheader();writer.writerows(flat)
from openpyxl import Workbook
wb=Workbook();ws=wb.active;ws.title='Milestone';ws.append(['Field','Value'])
for k,v in status.items():ws.append([k,json.dumps(v) if isinstance(v,(list,dict)) else v])
ws=wb.create_sheet('Performance');ws.append(list(flat[0]));[ws.append(list(r.values())) for r in flat]
ws=wb.create_sheet('Work normalization');ws.append(list(normalizations[0]));[ws.append(list(r.values())) for r in normalizations]
ws=wb.create_sheet('Failure contracts');ws.append(list(failureRows[0]));[ws.append(list(r.values())) for r in failureRows]
ws=wb.create_sheet('Inherited B8 issues')
with (E/'inherited-B8-issue-priorities.csv').open() as file:
 for row in csv.reader(file):ws.append(row)
for ws in wb:
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
wb.save(B/'milestone-status.xlsx')
print('B9B reports and local JSON/CSV/XLSX tracker rendered')
