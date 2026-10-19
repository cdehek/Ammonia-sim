"""Generate reports from qualified evidence, without editing historical research."""
import json,csv,hashlib,gzip,statistics,shutil
from pathlib import Path
from openpyxl import Workbook
B=Path(__file__).resolve().parent;E=B/'evidence';R=B.parents[2]
q=json.loads((E/'qualification.json').read_text());p=json.loads((E/'performance-summary.json').read_text());assert q['passed']
def write(n,s):(B/n).write_text(s.strip()+'\n')
def read(n):return json.loads(gzip.decompress((E/n).read_bytes()))
def table(headers,rows):return '| '+' | '.join(headers)+' |\n|'+ '|'.join(['---']*len(headers))+'|\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)
improved=all(v['improvementMeasured'] for e in p.values() for v in e['fixtures'].values())
perf='YES' if improved else 'INCONCLUSIVE'
decisions={'Source identity':'PASS','Table purity proof':'PASS','Cache ownership':'PASS','Scientific parity':'PASS','Failure/rollback':'PASS','Boundary/phase regression':'PASS','Property-call reduction measured':'YES','Performance improvement measured':perf,'Further isolated research':'GO','Operator integration':'NO-GO'}
status='B9C EXACT BOUNDARY PH REUSE PASS'
write('ENGINEERING_REPORT.md',f'''# B9C engineering disposition

**{status}**

The isolated original-table prototype preserves the source-proven contract and
every tested scientific/failure outcome. It removes only duplicate successful
boundary PH+SAT evaluations within a single actual nonlinear stage. It does not
change the five evaporator inventories, residual arithmetic, signed FD, Newton,
line search, TR-BDF2, adaptation, conservation or acceptance budgets.

{table(['Required separate decision','Result'],decisions.items())}

Measurable performance improvement: **{perf}**. Timing results below and in
BROWSER_PERFORMANCE_MATRIX.md are measured matched browser results, not a
prediction from the eliminated call count. Gate 4 is a reliable decision about
this bounded experiment, not production or sustained real-time qualification.

Gate 1 source review was committed independently before functional code at
36b118b30e44613f1c7ac9ca615aa61c949a88c2. It is a review independent of the
implementation, not third-party certification. Gate 2 compares full exact
science/observer evidence, every residual input/output hash, raw residual arrays,
Jacobians/corrections, converged stages and original saved B9B same-engine hashes.
Gate 3 covers actual boundary/phase/patch/thermal/numerical tests and unchanged
C1 fail-closed tests. Gate 4 uses five predeclared sequential C→candidate→C
brackets per fixture/engine, separate diagnostic passes and all samples retained.

Cold/warm real PH counts are 8,476→2,756 and 22,072→6,892 in BOTH browsers.
The 5,720/15,180 removed PH calls match the measured upper bound in these cases.
Residual work stays 2,914/7,710; stages 54/120; corrections 130/345; accepted
half-steps 16/34; accuracy rejections 1/3. No non-stage PH call is removed.
Every omitted PH has original exact p/h and the same first successful value;
the projected remaining public PH/SAT/MVU sequence is exactly unchanged.

Purity is conditional on the privately owned, recursively frozen original table
and recognized fixed engine. The general createEngine(data) API is mutable and
does NOT receive a blanket purity claim. Unowned/injected engines are refused.
No C1/CoolProp/WASM route uses the new cache. Failed results are never installed;
any residual exception clears both slots. There is no global/property-state cache.

Qualification limits: only bounded startup and extra bounded regressions; no
full A2 campaign, longer horizon, experimental backend repair, equipment
calibration, production protection or physical iPadOS test. Missing liquid-table
rows are tested through the original statePT interface; PH's original branchLimit
avoids those rows. A failed in-stage boundary cannot be installed through valid
create/update interfaces; exact-source mechanical tests cover failure/invalid-key
cache mechanics without adding an injected backend to the optimized model.
Thermal-before-boundary precedence is proven from the unchanged source order;
actual outward 200/400 K trajectories exercise the original stop/rollback path.

Next milestone: independent engineering review of the minimal diff and retained
receipts; if separately authorized, a small hardening study of the table-only
ownership API and broader supported boundary trajectories. Do not automatically
begin another optimization, integration, campaign, merge or deployment.
Operator workstation integration remains **NO-GO**. Passing B9C does not authorize
merging this prototype into accepted A2.
''')
prop=[];science=[];timing=[];overhead=[]
for engine,e in q['engines'].items():
 for fixture,v in e['fixtures'].items():
  d=v['property'];prop.append([engine,fixture,d['logicalStageBoundaryRequests']['inlet'],d['logicalStageBoundaryRequests']['outlet'],d['candidateMisses']['inlet'],d['candidateMisses']['outlet'],d['successfulHits'],d['realPHControl'],d['realPHCandidate'],d['nonStagePHControl']])
  w=v['work'];science.append([engine,fixture,'EXACT',w['residuals'],d['stages'],w['corrections'],w['accepted'],w['rejected']])
  r=p[engine]['fixtures'][fixture];a=r['values']['control']['advanceMS'];b=r['values']['candidate']['advanceMS'];timing.append([engine,fixture,f"{a['median']:.2f} [{a['min']:.2f}–{a['max']:.2f}]",f"{b['median']:.2f} [{b['min']:.2f}–{b['max']:.2f}]",f"{r['medianPairedReductionPercent']:.1f}%",r['decision']])
  for mode in ['control','candidate']:
   diag=r['diagnosticAdvanceMS'][mode];normal=r['values'][mode]['advanceMS']['median'];overhead.append([engine,fixture,mode,f'{normal:.2f}',f'{diag:.2f}',f'{diag/normal:.2f}×'])
write('PROPERTY_CALL_REDUCTION.md','''# Measured property-call reduction

'''+table(['Engine','Fixture','Logical inlet','Logical outlet','Inlet misses','Outlet misses','Hits / PH saved','All real PH control','All real PH candidate','Non-stage PH each'],prop)+'''

Logical requests remain one inlet and outlet per residual. In the fixed fixtures
each stage retains one successful call of each boundary. Savings include solved
stages in accuracy-rejected trials, whose slots are discarded before retry.
No reduction in residual work is counted as savings.

evidence/qualification.json records public SAT, nested lexical SAT, independent
canonical recoveries, section-cache requests/hits/misses, non-stage boundary
requests, outside-stage PH and thrown operations separately. The primary raw
original/candidate ledgers are evidence/parity-{chromium,webkit}.json.gz under
runs[].evidence.properties and .boundaries. All numeric fields are tagged binary64.
The property list is an actual public engine-call ledger; boundaries lists all
logical requests with realPH=0/1. Section hits are counted separately.

The analyzer checks every logical boundary p/h/result exactly, identifies each
stage's retained value, proves every omitted call equal to it, and removes ONLY
that request's original successful engine.ph followed by public engine.sat.
The remaining sequence must equal the complete candidate sequence, including
inputs, outputs, throws, stage/role/request identities and order. Nested lexical
saturation calls disappear as a consequence of the omitted genuine engine.ph,
not as a second optimization. Raw lexical counts corroborate this attribution.

Trial direction/initial rates, coordinates, recovery, domain proof, create/update
validation, event checks, records and outlet diagnostics remain genuine calls.
The full fixtures retain 476 cold / 973 warm outside-stage PH calls each.
Original thrown property operations are zero in the fixed fixtures. Refusal
regressions retain their original errors and are reported separately.
''')
write('SCIENTIFIC_PARITY_MATRIX.md','''# Exact scientific parity

'''+table(['Engine','Fixture','Result','Residuals','Stages','Newton corrections','Accepted half-steps','Accuracy rejects'],science)+'''

PASS against both the isolated unchanged control and saved B9B SAME-engine
full scientific hashes at 0/.0001/.0005/.001 seconds. Requested advances remain
.0001/.0004/.0005 with original config/settings and initial, mandatory and extra
snapshot records. No easier performance fixture is substituted.

The full state/record/attempt encoding covers all authoritative M/U/V, guesses,
recovered p/h/T/rho/phase/quality/wet, thermal energies/temperatures, heat and flow
rates, mass/enthalpy fluxes, conservative residuals, ledgers, events/timestamps,
TR-BDF2 errors, acceptance/rejections, next proposals and final accepted states.
Every residual callback's complete return (including inventories/rates/quadrature)
and coordinates are SHA-256 encoded before optional detail reduction. Raw x and
residual arrays, full matrices/rhs/Newton delta, and full stage results remain.
Original observer events retain every signed FD, line search, iteration and
stage context. Independent hashes and raw component equality both pass.

All original converged residuals and section/aggregate conservation budgets pass.
auditRecord independently recomputes signed thermal ledgers. Canonical recovery
and thermal closure use the original unchanged comparisons and source budgets.
Every original internal adaptive attempt/error and endpoint proposal is exact;
unchanged Time.factor plus identical errors preserves the complete proposal
sequence. The preserved B9B same-browser adaptive-proposal-reconstruction.json
already checks these identical captured inputs and endpoints independently.

The frozen startup fixtures have no natural phase events; their empty event lists
are compared exactly. Boundary phase cases do not claim complete inventory phase
crossing or topology qualification.

No tiny numerical tolerance replaces bit equality. Cross-browser values are
not assumed bit-identical; each browser matches its own original baseline.
evidence/qualification.json includes exact hashes and conservation receipts.
''')
regrows=[]
for name in q['engines']['chromium']['regressions']:regrows.append([name,'EXACT in Chromium and WebKit'])
write('BOUNDARY_AND_PHASE_REGRESSIONS.md','''# Boundary and phase regressions

'''+table(['Bounded original-interface scenario','Result'],regrows)+'''

In addition to frozen cold/warm, each engine executes ten bounded scenario pairs,
31 PH input/phase cases, four refused boundary updates and an original missing-row
statePT refusal. The raw cases and bit encodings remain in regressions-*.json.gz.
All full states, records, attempts, component hashes and logical boundary-result
hashes agree. Property counts differ only by successful stage-boundary hits.

Reversal changes inlet 4→3 bar and outlet 3.5→4 bar after initial supported work.
Both end flows reverse. The outlet now donates 1700 kJ/kg; energyKW is checked
against the signed mass flow times that actual boundary enthalpy. Sections retain
their own dynamic state. Scheduled and manual patches use original paths;
the scheduled event occurs exactly at .0001 s and later stages see 4.2/3.4 bar.
No prior stage cache is reachable from either patch path.

Saturated liquid, vapor and genuine x=.4 mixture boundaries use original sat
values; genuine superheated outlet uses the original expensive inverse. Direct
hf/hg near-edge inputs preserve the original eps classification. Pressures test
.3±1e-12 and 35±1e-12 bar, endpoints, nonfinite and zero inputs. Enthalpies include
supported/subcool/superheat edges, nearby supported and refused states, signed
zero and nonfinite inputs. No epsilon support or physical input change is added.

Cold/warm naturally reject 1/3 accuracy steps. All retries preserve exact original
decisions and successful candidates never contaminate accepted inventories.
The bounded numerical-stop fixture deliberately selects supported restrictive
maxIterations=1,maxAttempts=20 to exercise a numerical stop; these settings NEVER
enter the fixed cold/warm qualification or timing workloads.

Limits: PH branchLimit avoids absent liquid rows, so missing-row refusal is tested
through original statePT, not forced by editing the solver. A caller cannot install
an unsupported boundary using valid create/update. In-stage failure cache behavior
is verified by an exact-source isolated mechanical unit, not a new backend route.
Thermal validation precedence is source-proven; supported genuine outward 200/400 K
cases exercise refusal/rollback without forging invalid accepted state.
''')
write('FAILURE_AND_ROLLBACK.md','''# Failure, rollback and lifecycle integrity

PASS in actual Chromium and desktop WebKit. C1 source, adapter, patches and
scientific artifacts are Git-verified unchanged. c1-*.json.gz reruns the original
B9 local C1 suite plus unchanged B9A c1Session transient, persistent, boundary,
failed-record and genuine thermal200/thermal400 scenarios. No C1 source imports
or selects the B9C optimized model; only the original control table is supplied.

The original assertions cover supported five-section work, transient synthetic
refusal, persistent synthetic refusal, numerical stop, 20 attempts / 19 exact
halvings, minStep=1e-9 s, maxIterations=30, per-rejection binary64 M/U/V/thermal/
ledger/clock/event rollback and object identity, accepted-state record refusal,
retained accepted clock/events and independently confirmed outward thermal exits.
Inverse refusals remain numerical/unproved; no artificial physical trip or safe
dry-suction state is invented. Synthetic EOS evidence is explicitly synthetic.

For the table-only prototype, strict invalid updates leave the entire state
bit-exact; genuine 200/400 K stops preserve physical bits and references; numerical
stop and accuracy-retry decisions match control. The mechanical cache unit extracts
the ACTUAL prototype closure and catch expression, checks exact changed inputs,
signed zero, NaN, repeated failure, original Error identity, no failed installation
and invalidation of BOTH slots after failure. It does not inject a model backend.
The candidate guard refuses an unowned table engine and therefore any CoolProp
adapter. Engine/table freezing and source scope are independently verified.

Every request uses a fresh actual worker, terminated after response/error.
All timed fresh-worker initial states and final scientific hashes match, proving
restart equivalence in the tested sessions. There is no global boundary cache,
accepted-state property pointer or cross-session cache transport. Memory observations
are bounded high-water evidence, not a proof of general leak freedom. Existing
broader B9/B9A protocol tests remain protected evidence; this experiment does not
claim a new production protocol or UI qualification.
''')
write('BROWSER_PERFORMANCE_MATRIX.md','''# Matched actual-browser performance

'''+table(['Engine','Fixture','Control ms n=10, median [range]','Candidate ms n=5, median [range]','Median paired reduction','Decision'],timing)+'''

Times are original three-advance sums to .001 simulated seconds, including all
retry/stage/commit and mandatory record work. Extra snapshot records, cloning,
hashing, JSON artifact export, communication, browser/worker startup and table
ownership construction are outside those advance timers. Original observer cadence
remains active. No heavy diagnostic wrappers/counters run in primary timing.

Five predeclared C→candidate→C brackets per fixture/engine; fixture order reverses
on alternating rounds. Workloads are strictly sequential, with fresh workers.
All samples and before/after variability remain in timing-*.json.gz and
performance-summary.json. Every timed scientific hash matches qualified control.
No outlier is removed. Ranges, sample standard deviations, paired deltas and raw
samples are reported; these sparse rounds do not establish population confidence
intervals or sustained throughput. A consistent bracket reduction is the stated
bounded improvement criterion, not an inferred speedup from call counts.

Construction and model create+initial-record times, per-advance values, extra
record generation and full request RTT are separate in the raw samples. RTT
includes research module loading, full evidence hashing and serialization; it is
not pure worker/main communication cost. A pure communication-only duration is
not isolated here. Stage/property/residual timings are separate DIAGNOSTIC passes
with explicitly measured overhead, never added as an exclusive CPU partition.

Chromium and desktop WPE WebKit versions, executable identities, user agents and
source hashes are recorded in raw receipts. No physical iPadOS Safari inference.
Main-page 10-ms heartbeat is measured during primary research requests. The minimal
harness has no operator input UI; real UI input responsiveness is NOT APPLICABLE.
Available page JS heap and whole-browser RSS observations are retained; WebKit
may lack performance.memory. RSS double-counts shared pages and does not identify
live cache bytes. All workers terminate and fresh sessions restart exactly.
Two boundary result objects per active stage are the only new retained values;
no evidence establishes long-duration leak freedom or operator readiness.
''')
write('INSTRUMENTATION_OVERHEAD.md','''# Diagnostic overhead and limits

'''+table(['Engine','Fixture','Mode','Primary median ms','Diagnostic ms (one pass)','Ratio'],overhead)+'''

The diagnostic path instruments actual public PH/SAT/MVU and lexical saturation/
inverse/branch operations, stage/trial/rate/record scopes, every residual result,
Jacobian/correction and logical boundary request. It stores values already
computed, adding no scientific property/residual calls. Timers and output data
are detached from numerical decisions. Function bodies are separately labeled
diagnostic artifacts with exact patches and hashes.

Diagnostic timers retain inclusive stage, residual, property-engine, rate and
record intervals. They overlap and include bookkeeping; they are not additive
or intrinsic exclusive CPU costs. Tiny intervals quantize at browser resolution.
The full residual-value SHA is calculated AFTER advance, before reduction to
raw residual arrays. Matrix/FD/Newton/stage equality is independent of timing.

Ratios compare one diagnostic run with primary matched medians, not a stable
overhead estimate or paired confidence interval. The initial overly detailed
artifact exports were manually terminated and superseded: they imposed excessive
protocol/serialization memory/time. They produced no scientific failure or usable
performance result. Final exports preserve complete fixed-workload property
ledgers and science, with full residual-value hashes plus raw residual/matrix/
stage evidence. Extra regressions retain complete science/observer data, component
hashes and logical boundary-result hashes with counts instead of duplicating all
large raw diagnostic arrays. Failed/superseded logs are retained explicitly.
No favorable timing sample selection followed those harness-only corrections.
''')
write('SOURCE_DIFF_AND_HASH_AUDIT.md','''# Source identity and minimal-diff review

PASS. Gate 0 hashes all 2,791 historical tracked files; fourteen accepted source
bodies match exact A2 Git objects. Later root coupled-model/Newton are excluded.
Control-bundle contains unchanged source bodies. Candidate changes only the
B9C-owned isolated/coupled-model.cjs and adds a table-only ownership wrapper.
No original accepted snapshot, property source, solver, UI, Live Plant or hosting
file is edited. The prototype is not a replacement accepted A2 implementation.

BOUNDARY_PH_PROTOTYPE.patch is the complete functional source diff. It adds the
owned-engine guard, optional private evaluate callback, two stage-local slots,
Object.is exact checks, lazy original thermo.ph, residual-exception clear and one
endpoint plumbing argument. All original arithmetic expressions and evaluation
order remain. verify-source.py independently removes ONLY these declared changes
and restores the exact accepted bytes, then verifies the other thirteen bodies.
table-only.cjs recursively freezes the actual private table and freezes/weakly
registers engines. This guard excludes injected/mutable/experimental engines.

prototype-source-manifest.json identifies all source bodies and four generated
bundles. diagnostic-{control,candidate}.patch shows the separate observations;
none enters primary timing. Runtime receipts identify browser engines/launchers;
final runtime-manifest.json additionally hashes installed browser binaries and
all delivered research sources. Original saved B9/B9A/B9B/C1/B6/B7 artifacts are
Git-verified, not rebuilt or requalified. The final preservation receipt checks
every historical file, full starting remote map and PR #3 identities/metadata.
''')
write('MECHANICS_REPORT.md',f'''# Mechanic’s report

The computer repeatedly works out temperature, density, phase and wet/superheat
indications from the imposed inlet and outlet pressure and enthalpy. Even a table
lookup can require a bounded inverse calculation, especially at the superheated
outlet. Newton asks for many slightly perturbed section states while finding one
acceptable numerical stage, but those two imposed boundary conditions stay fixed.

This prototype retains each boundary's first successful result for that one stage.
When inlet/outlet pressure or enthalpy changes, the original validation and property
calculation run again. New stages, retries, patches and workers have new slots.
It does not freeze the refrigerant in the five evaporator sections: their mass,
energy, pressure, enthalpy and heat/flow responses continue through the original
finite-volume solver and its original checks.

Real property work fell by 5,720 cold / 15,180 warm PH evaluations, in both tested
browsers. All tested refrigerant inventories, thermal energies, flows, temperatures,
solver decisions and conservation results stayed bit-exact. Numerical refusals
remain refusals with the original rollback; genuine outward 200/400 K stops keep
their original meaning. Nothing invents a compressor trip or dry-suction assurance.

Measured performance improvement: **{perf}** for these tiny startup workloads;
the actual matched times and variability are in BROWSER_PERFORMANCE_MATRIX.md.
This does not establish sustained real-time operation or physical-device speed.
CoolProp has a different implementation, ownership and failure contract, so the
table purity proof cannot be transferred to it. No CoolProp repair was performed.
This research harness is **not ready for operator use**. Operator integration
remains NO-GO, and the prototype requires engineering review before any separately
authorized next step. It is not merged or deployed.
''')
write('README.md',f'''# Research B9C — exact stage-local boundary PH reuse

**{status}**. Performance improvement measured: **{perf}**.
Original table only; isolated prototype; operator integration **NO-GO**.
No merge, deployment, production UI, Live Plant, accepted A2 or EOS change.

Read ENGINEERING_REPORT.md, BOUNDARY_PH_EQUIVALENCE_DESIGN.md,
SOURCE_DIFF_AND_HASH_AUDIT.md and the scientific/failure/performance matrices.
The independent pre-implementation design commit is
36b118b30e44613f1c7ac9ca615aa61c949a88c2, based on exact B9B
ec2df3f9ee3efd7faef3d129282ad1dd7bdc896f. The canonical C2 correction inherited
from B8 is recorded in evidence/gate0.json; no protected ref was changed.

## Reproduction

Use Python 3, Playwright 1.62.0, openpyxl 3.1.5, Node 24 and the pinned browser setup. Network is
needed only for repository/browser setup; the harness serves local files. Read
the managed-environment proxy instructions before network setup. No EOS rebuild.
Run from repository root, writing NEW receipt names/directories when reproducing
so that committed evidence is not overwritten. Each browser command has an outer
bounded timeout; per-worker timeout is 90 s.

```sh
python3 experiments/dx-fv/browser_eos_b9c/build.py
python3 experiments/dx-fv/browser_eos_b9c/verify-source.py
timeout 420 python3 experiments/dx-fv/browser_eos_b9/setup-browsers.py
export PLAYWRIGHT_BROWSERS_PATH=/workspace/b9-runtime/browsers
timeout 300 python3 experiments/dx-fv/browser_eos_b9c/run-browser.py --stage parity
timeout 300 python3 experiments/dx-fv/browser_eos_b9c/run-browser.py --stage regressions
timeout 180 python3 experiments/dx-fv/browser_eos_b9c/run-browser.py --stage c1
timeout 90 python3 experiments/dx-fv/browser_eos_b9c/analyze.py
timeout 300 python3 experiments/dx-fv/browser_eos_b9c/run-browser.py --stage timing
python3 experiments/dx-fv/browser_eos_b9c/analyze.py --performance
python3 experiments/dx-fv/browser_eos_b9c/write-reports.py
```

run-browser refuses existing stage receipts. For a new run, use a fresh isolated
worktree and archive the committed B9C stage receipts into a separate B9C-owned
subdirectory before running; never alter historical B9/B9A/B9B evidence. The
immutable original baseline remains read-only. gate0.py is intentionally restricted
to the starting B9B HEAD; final preservation.py uses the starting SHA and hashes.

Evidence encoding uses f64 hexadecimal binary64 tags, including NaN/signed zero;
equality never converts them to toleranced decimal values. Primary performance
samples contain exact science hashes, settings and original work counters. Raw
diagnostic ledgers contain the actual property requests and results. Extra
regressions retain hashes of complete residual/linear/stage data and full original
state/records/observer events. Source manifests and diagnostic patches make those
different evidence sizes explicit. No phase/physics acceptance is inferred from
timing or a compact property vector.

milestone-status.csv/json/xlsx are local tracker-compatible artifacts. No external
tracker, PR comment or message is posted. Protected refs and PR #3 receipts are
local research evidence. Stop for engineering review after the authorized branch
push and final clean-tree/preservation checks.
''')
shutil.copyfile(B.parent/'browser_eos_b9b/SCIENTIFIC_CONTRACTS_AND_BLOCKERS.md',B/'PRESERVED_SCIENTIFIC_CONTRACTS_AND_BLOCKERS.md')
shutil.copyfile(B.parent/'browser_eos_b9b/evidence/inherited-B8-issue-priorities.json',E/'inherited-B8-issue-priorities.json')
shutil.copyfile(B.parent/'browser_eos_b9b/evidence/inherited-B8-issue-priorities.csv',E/'inherited-B8-issue-priorities.csv')
rows=[{'milestone':'B9C','decision':k,'status':v,'disposition':status} for k,v in decisions.items()]
(B/'milestone-status.json').write_text(json.dumps(rows,indent=2)+'\n')
with (B/'milestone-status.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
wb=Workbook();ws=wb.active;ws.title='B9C';ws.append(list(rows[0]));[ws.append(list(r.values())) for r in rows];ws.freeze_panes='A2'
ws=wb.create_sheet('Timing samples');ws.append(['Engine','Fixture','Round','Before ms','Candidate ms','After ms','Reduction %'])
for engine,e in p.items():
 for fixture,v in e['fixtures'].items():
  for r in v['paired']:ws.append([engine,fixture,r['round'],r['beforeMS'],r['candidateMS'],r['afterMS'],r['reductionPercent']])
ws=wb.create_sheet('Inherited blockers');block=json.loads((E/'inherited-B8-issue-priorities.json').read_text());ws.append(['Preserved exact JSON']);ws.append([json.dumps(block)[:32767]])
wb.save(B/'milestone-status.xlsx')
component=[]
for engine,e in p.items():
 for fixture,v in e['fixtures'].items():
  for mode,t in v['diagnosticScopesMS'].items():component.append([engine,fixture,mode]+[f'{t.get(k,0):.2f}' for k in ['stage','residual','property.ph','property.sat','property.stateMVU','record']])
with (B/'INSTRUMENTATION_OVERHEAD.md').open('a') as f:
 f.write('\n'+table(['Engine','Fixture','Mode','Stage ms','Residual ms','PH ms','Public SAT ms','MVU ms','Record ms'],component)+'\n\nStage/residual timers cover advance work. Property and record scopes cover the\nfull frozen workload, including initialization and extra records. These are\none-pass inclusive DIAGNOSTIC measurements with overhead, not primary medians.\n')
with (E/'matched-performance-samples.csv').open('w',newline='') as f:
 w=csv.writer(f);w.writerow(['engine','fixture','round','control_before_ms','candidate_ms','control_after_ms','paired_reduction_percent'])
 for engine,e in p.items():
  for fixture,v in e['fixtures'].items():
   for r in v['paired']:w.writerow([engine,fixture,r['round'],r['beforeMS'],r['candidateMS'],r['afterMS'],r['reductionPercent']])
index={engine:{fixture:{'artifact':f'parity-{engine}.json.gz','control':'runs[fixture='+fixture+',mode=control].evidence','candidate':'runs[fixture='+fixture+',mode=candidate].evidence','scienceSHA256':v['scienceSHA256'],'actualCounts':v['property']} for fixture,v in e['fixtures'].items()} for engine,e in q['engines'].items()}
(E/'property-call-ledger-index.json').write_text(json.dumps(index,indent=2)+'\n')
print(status,'; performance improvement:',perf)
